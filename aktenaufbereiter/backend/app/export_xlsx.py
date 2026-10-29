from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .schemas import Case
from .tables import build


def export(case: Case, ann: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for t in build(case, ann).values():
        ws = wb.create_sheet(t["title"][:31].replace(":", ""))
        ws.append(t["cols"])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="EAEAEA")
        for r in t["rows"]:
            ws.append(r["cells"])
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[col[0].column_letter].width = 14 if i == 1 else 34
            for c in col:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
